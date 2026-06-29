from pathlib import Path
import fitz
from docx import Document

def generate_samples():
    docs_dir = Path("data/docs")
    docs_dir.mkdir(parents=True, exist_ok=True)

    # 1. Create DOCX Password Policy
    doc = Document()
    doc.add_heading("Company Password Policy", level=1)
    doc.add_paragraph("Employees must use passwords of at least 14 characters.")
    doc.add_paragraph("Passwords must be changed every 90 days to prevent compromised credential risks.")
    doc.add_paragraph("Passwords must not be shared externally or stored in plaintext files.")
    doc.add_paragraph("All accounts with administrative access require multi-factor authentication (MFA).")
    doc.save(docs_dir / "example_password_policy.docx")
    print("Generated: example_password_policy.docx")

    # 2. Create DOCX Incident Reporting Policy
    doc2 = Document()
    doc2.add_heading("Security Incident Reporting Policy", level=1)
    doc2.add_paragraph("Security incidents must be reported to the SOC within 1 hour of discovery.")
    doc2.add_paragraph("Any suspected data breach must be escalated immediately to the Information Security team and Legal.")
    doc2.add_paragraph("Employees must preserve evidence and avoid unauthorized remediation or communication about the incident.")
    doc2.save(docs_dir / "example_incident_reporting.docx")
    print("Generated: example_incident_reporting.docx")

    # 3. Create PDF Remote Work Controls
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((50, 50), "Remote Work Security Controls Policy", fontsize=16, fontname="hebo")
    
    controls = [
        "1. Remote employees must use company-approved and managed devices only.",
        "2. Connecting via public Wi-Fi is prohibited unless a corporate VPN is active.",
        "3. Customer data and corporate files may not be stored on personal devices or cloud drives.",
        "4. Work screens must be locked immediately when left unattended in remote settings.",
        "5. Compliance reviews are conducted quarterly on all remote access logs."
    ]
    
    y = 90
    for control in controls:
        page.insert_text((50, y), control, fontsize=11, fontname="helv")
        y += 25
        
    pdf.save(docs_dir / "example_remote_work.pdf")
    pdf.close()
    print("Generated: example_remote_work.pdf")

if __name__ == "__main__":
    generate_samples()
